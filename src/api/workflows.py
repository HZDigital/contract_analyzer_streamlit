"""Synchronous, storage-agnostic adapters for every supported analysis workflow.

The jobs layer supplies the uploaded bytes and persists the returned artifacts. This
module deliberately has no knowledge of HTTP requests, UI components, or Blob storage.
"""

from __future__ import annotations

import csv
import json
import logging
import math
import re
from collections.abc import Callable, Iterable, Mapping
from dataclasses import asdict, dataclass, is_dataclass
from datetime import date, datetime
from enum import Enum
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.utils.ai_analyzer import (
    analyze_contract,
    analyze_tender_document,
    analyze_tender_with_fields,
    compare_contracts,
    compare_factory_documents,
    extract_client_and_products,
    extract_client_and_products_from_invoices,
    group_similar_products,
)
from src.utils.invoice_normalstunden_extractor import extract_normalstunden_from_bytes
from src.utils.large_contract_analyzer import (
    analyze_contract_chunks,
    build_contract_chunks,
    merge_chunk_findings,
    synthesize_contract_analysis,
)
from src.utils.pdf_processor import (
    extract_pdf_pages,
    extract_text_from_file,
    extract_text_from_pdf,
)
from src.utils.web_research import analyze_market_situation

from .public_results import curate_public_result
from .settings import get_settings
from .uploads import expand_normalstunden_archive


ProgressCallback = Callable[[int, str], None]
logger = logging.getLogger(__name__)
STANDARD_CONTRACT_CHARACTER_LIMIT = 30_000
ANALYSIS_ERROR_MESSAGES = {
    "configuration_missing": "The analysis service is not configured.",
    "client_initialization_failed": "The analysis service configuration is invalid.",
    "request_failed": "The analysis service request failed. Please try again.",
    "response_invalid": "The analysis service returned an invalid response. Please try again.",
}


@dataclass(frozen=True)
class WorkflowInput:
    """A validated upload represented entirely in memory."""

    name: str
    role: str
    data: bytes
    content_type: str | None = None
    supplier_hint: str = ""


@dataclass(frozen=True)
class WorkflowArtifact:
    """An in-memory export for the jobs layer to persist as an artifact."""

    data: bytes
    name: str
    content_type: str


@dataclass(frozen=True)
class WorkflowResult:
    """Normalized result data plus exports produced by a workflow."""

    result: dict[str, Any]
    artifacts: list[WorkflowArtifact]

    @property
    def result_json(self) -> dict[str, Any]:
        """Alias describing the JSON-safe result persisted on a job record."""
        return self.result


def run_workflow(
    job: Any,
    input_files: Iterable[WorkflowInput | Mapping[str, Any] | Any],
    progress_callback: ProgressCallback | None = None,
) -> WorkflowResult:
    """Run one supported workflow synchronously using only supplied upload bytes.

    ``job`` is normally a ``JobRecord``. A mapping with ``workflow`` and optional
    ``options`` is also accepted so the adapter remains independent of persistence.
    Input values can be ``WorkflowInput``, ``ValidatedUpload``, or mappings with
    ``name``, ``role``, and ``data`` keys.
    """
    workflow = str(_job_value(job, "workflow", "")).strip()
    options = _as_mapping(_job_value(job, "options", {}))
    files = [_coerce_input_file(item) for item in input_files]
    if not files:
        raise ValueError("A workflow requires at least one input file.")

    report_progress = _progress_reporter(progress_callback)
    report_progress(0, "Starting workflow")

    handlers: dict[str, Callable[[list[WorkflowInput], Mapping[str, Any], ProgressCallback], WorkflowResult]] = {
        "product_request": _run_product_request,
        "invoice": _run_invoice,
        "normalstunden": _run_normalstunden,
        "detailed_contract": _run_detailed_contract,
        "tender": _run_tender,
        "cooperation_review": _run_cooperation_review,
        "factory_certificate": _run_factory_certificate,
        "large_scanner": _run_large_scanner,
    }
    handler = handlers.get(workflow)
    if handler is None:
        raise ValueError(f"Unsupported workflow: {workflow}")

    outcome = handler(files, options, report_progress)
    result = _normalize_json(outcome.result)
    if not isinstance(result, dict):
        raise TypeError("Workflow result must be a JSON object.")
    report_progress(100, "Workflow complete")
    return WorkflowResult(result=result, artifacts=outcome.artifacts)


def _run_product_request(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    results: list[dict[str, Any]] = []
    total = len(files)
    for index, file in enumerate(files, 1):
        progress(_file_progress(index - 1, total, 5, 70), f"Extracting {file.name}")
        try:
            text = extract_text_from_pdf(file.data)
            if not text.strip():
                raise ValueError("No readable text found in PDF")
            extracted = _json_mapping(extract_client_and_products(text))
            if extracted.get("error"):
                results.append(
                    {
                        "file_name": file.name,
                        "status": "failed",
                        "products": [],
                        "error": _analysis_error_message(extracted),
                    }
                )
            else:
                results.append(
                    {
                        "file_name": file.name,
                        "status": "success",
                        "client_name": extracted.get("client_name", "Not detected"),
                        "products": [
                            _select_fields(product, "product_name", "quantity", "unit", "description")
                            for product in _dict_list(extracted.get("products"))
                        ],
                        "contract_type": extracted.get("contract_type", "Unknown"),
                        "total_estimated_value": extracted.get("total_estimated_value", "Not specified"),
                    }
                )
        except Exception as exc:  # noqa: BLE001
            results.append({"file_name": file.name, "status": "failed", "products": [], "error": _safe_document_error(exc)})
        progress(_file_progress(index, total, 5, 70), f"Processed {file.name}")

    progress(75, "Grouping matching products")
    products = [
        {
            "product_name": product.get("product_name", "Unknown"),
            "quantity": product.get("quantity", "Not specified"),
            "unit": product.get("unit", ""),
            "client_name": result.get("client_name", "Not detected"),
            "contract_type": result.get("contract_type", "Unknown"),
            "description": product.get("description", "Not specified"),
        }
        for result in results
        if result.get("status") == "success"
        for product in _dict_list(result.get("products"))
    ]
    should_group = _option_bool(options, "groupSimilarProducts", "group_similar_products", default=True)
    grouping = _json_mapping(group_similar_products(products)) if products and should_group else {"groups": []}
    consolidated = _consolidate_products(products, _dict_list(grouping.get("groups"))) if should_group else []
    warnings = []
    if grouping.get("error"):
        warnings.append(
            "Products were extracted, but similar products could not be consolidated. Review each document result separately."
        )
    payload = {
        "workflow": "product_request",
        "summary": _summary(results),
        "results": results,
        "consolidated_products": consolidated,
        "warnings": warnings,
    }
    artifacts = [
        _json_artifact("product_request_results.json", payload),
        WorkflowArtifact(
            data=_csv_bytes(_product_export_rows(results)),
            name="product_request_results.csv",
            content_type="text/csv; charset=utf-8",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_invoice(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    del options
    results: list[dict[str, Any]] = []
    total = len(files)
    for index, file in enumerate(files, 1):
        progress(_file_progress(index - 1, total, 5, 80), f"Extracting {file.name}")
        try:
            text = extract_text_from_pdf(file.data)
            if not text.strip():
                raise ValueError("No readable text found in PDF")
            extracted = _json_mapping(extract_client_and_products_from_invoices(text))
            if extracted.get("error"):
                results.append(
                    {
                        "file_name": file.name,
                        "status": "failed",
                        "products": [],
                        "error": _analysis_error_message(extracted),
                    }
                )
            else:
                results.append(
                    {
                        "file_name": file.name,
                        "status": "success",
                        **_invoice_fields(extracted),
                    }
                )
        except Exception as exc:  # noqa: BLE001
            results.append({"file_name": file.name, "status": "failed", "products": [], "error": _safe_document_error(exc)})
        progress(_file_progress(index, total, 5, 80), f"Processed {file.name}")

    rows = _invoice_export_rows(results)
    payload = {"workflow": "invoice", "summary": _summary(results), "results": results}
    artifacts = [
        _json_artifact("invoice_results.json", payload),
        WorkflowArtifact(_csv_bytes(rows), "invoice_results.csv", "text/csv; charset=utf-8"),
        WorkflowArtifact(
            _xlsx_bytes(rows, "Invoices"),
            "invoice_results.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_normalstunden(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    pdf_files = _normalstunden_pdf_inputs(files, options)
    hints = _as_mapping(options.get("supplierHints", options.get("supplier_hints", {})))
    default_hint = str(options.get("supplierHint", options.get("supplier_hint", ""))).strip()
    results: list[dict[str, Any]] = []
    total = len(pdf_files)
    for index, file in enumerate(pdf_files, 1):
        progress(_file_progress(index - 1, total, 5, 85), f"Extracting Normalstunden from {file.name}")
        supplier_hint = str(hints.get(file.name) or file.supplier_hint or default_hint).strip()
        try:
            results.append(extract_normalstunden_from_bytes(file.data, file.name, supplier_hint))
        except Exception as exc:  # noqa: BLE001
            results.append(
                {
                    "file_name": file.name,
                    "file_path": "",
                    "supplier": supplier_hint,
                    "supplier_folder": supplier_hint,
                    "hours_total": None,
                    "hourly_rates": [],
                    "hourly_rate_display": "",
                    "entries": [],
                    "status": "failed",
                    "matched_rows": 0,
                    "error": _safe_document_error(exc),
                }
            )
        progress(_file_progress(index, total, 5, 85), f"Processed {file.name}")

    rows = _normalstunden_export_rows(results)
    payload = {"workflow": "normalstunden", "summary": _summary(results), "results": results}
    artifacts = [
        _json_artifact("normalstunden_results.json", payload),
        WorkflowArtifact(_csv_bytes(rows), "normalstunden_results.csv", "text/csv; charset=utf-8"),
        WorkflowArtifact(
            _xlsx_bytes(rows, "Normalstunden"),
            "normalstunden_results.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_detailed_contract(
    files: list[WorkflowInput], _options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    results: list[dict[str, Any]] = []
    total = len(files)
    for index, file in enumerate(files, 1):
        progress(_file_progress(index - 1, total, 5, 85), f"Extracting {file.name}")
        try:
            text = extract_text_from_pdf(file.data)
            if not text.strip():
                raise ValueError("No readable text found in PDF")
            truncate_length = min(STANDARD_CONTRACT_CHARACTER_LIMIT, len(text))
            analysis = _json_mapping(analyze_contract(text, max(1, truncate_length)))
            if analysis.get("error"):
                results.append(
                    {
                        "file_name": file.name,
                        "status": "failed",
                        "error": _analysis_error_message(analysis),
                        "analysis": {},
                    }
                )
            else:
                results.append(
                    {
                        "file_name": file.name,
                        "status": "success",
                        "analysis": _curated_analysis("detailed_contract", analysis),
                    }
                )
        except Exception as exc:  # noqa: BLE001
            logger.warning(
                "Document processing failed: workflow=detailed_contract error_type=%s",
                type(exc).__name__,
            )
            results.append({"file_name": file.name, "status": "failed", "error": _safe_document_error(exc), "analysis": {}})
        progress(_file_progress(index, total, 5, 85), f"Analyzed {file.name}")

    payload = {"workflow": "detailed_contract", "summary": _summary(results), "results": results}
    artifacts = [
        _json_artifact("detailed_contract_results.json", payload),
        WorkflowArtifact(
            _csv_bytes(_detailed_export_rows(results)),
            "detailed_contract_results.csv",
            "text/csv; charset=utf-8",
        ),
        WorkflowArtifact(
            _detailed_markdown(results).encode("utf-8"),
            "detailed_contract_report.md",
            "text/markdown; charset=utf-8",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_tender(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    tender_files = _files_for_role(files, "tenderDocuments")
    template_files = _files_for_role(files, "tenderTemplate")
    if not tender_files or len(template_files) != 1:
        raise ValueError("Tender workflow requires tender documents and one XLSX template.")

    template = template_files[0]
    workbook, worksheet, field_cells, source_cell = _load_tender_template(template.data)
    desired_fields = list(field_cells)
    display_fields = _tender_display_fields(desired_fields)
    results: list[dict[str, Any]] = []
    total = len(tender_files)
    market_research = _option_bool(options, "includeMarketResearch", "include_market_research", default=True)
    for index, file in enumerate(tender_files, 1):
        progress(_file_progress(index - 1, total, 10, 75), f"Analyzing tender {file.name}")
        try:
            text = extract_text_from_pdf(file.data)
            if not text.strip():
                raise ValueError("No readable text found in PDF")
            analysis = _json_mapping(
                analyze_tender_with_fields(text, desired_fields) if desired_fields else analyze_tender_document(text)
            )
            if analysis.get("error"):
                results.append(
                    {
                        "file_name": file.name,
                        "status": "failed",
                        "error": _analysis_error_message(analysis),
                        "analysis": {},
                    }
                )
            else:
                _limit_tender_fields(analysis, display_fields, template_fields=bool(desired_fields))
                if market_research:
                    _add_market_situation(analysis)
                results.append(
                    {
                        "file_name": file.name,
                        "status": "success",
                        "analysis": _curated_analysis("tender", analysis, display_fields=display_fields),
                    }
                )
        except Exception as exc:  # noqa: BLE001
            results.append({"file_name": file.name, "status": "failed", "error": _safe_document_error(exc), "analysis": {}})
        progress(_file_progress(index, total, 10, 75), f"Processed tender {file.name}")

    progress(80, "Filling tender template")
    merged = _merge_tender_results(results)
    _fill_tender_template(worksheet, field_cells, source_cell, merged, results)
    if not field_cells:
        _add_tender_results_sheet(workbook, results)
    workbook_output = BytesIO()
    try:
        workbook.save(workbook_output)
    finally:
        workbook.close()

    payload = {
        "workflow": "tender",
        "summary": _summary(results),
        "template_name": template.name,
        "_display_fields": display_fields,
        "results": results,
        "merged": merged,
    }
    artifacts = [
        _json_artifact("tender_results.json", payload),
        WorkflowArtifact(
            workbook_output.getvalue(),
            "tender_template_filled.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        WorkflowArtifact(
            _tender_markdown(merged).encode("utf-8"),
            "tender_summary.md",
            "text/markdown; charset=utf-8",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_cooperation_review(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    suppliers = _files_for_role(files, "supplierAgreements")
    standards = _files_for_role(files, "standardContract")
    if not suppliers or len(standards) != 1:
        raise ValueError("Cooperation review requires supplier agreements and one standard contract.")

    progress(5, f"Extracting standard contract {standards[0].name}")
    standard_text = extract_text_from_file(standards[0].data, standards[0].name)
    supplier_texts: list[str] = []
    total = len(suppliers)
    for index, file in enumerate(suppliers, 1):
        progress(_file_progress(index - 1, total, 10, 55), f"Extracting {file.name}")
        supplier_texts.append(f"=== {file.name} ===\n{extract_text_from_file(file.data, file.name)}")
        progress(_file_progress(index, total, 10, 55), f"Processed {file.name}")

    progress(60, "Comparing cooperation agreements")
    analysis = _json_mapping(
        compare_contracts(
            "\n\n--- Document Separator ---\n\n".join(supplier_texts),
            standard_text,
            truncate_length=_optional_positive_int(options.get("truncateLength", options.get("truncate_length"))) or 12000,
            include_risk_assessment=_option_bool(options, "includeRiskAssessment", "include_risk_assessment", default=True),
            include_deviation_analysis=_option_bool(options, "includeDeviationAnalysis", "include_deviation_analysis", default=True),
            include_recommendations=_option_bool(options, "includeRecommendations", "include_recommendations", default=True),
        )
    )
    if analysis.get("error"):
        raise RuntimeError("The analysis service could not compare the supplied agreements.")
    analysis = _curated_analysis("cooperation_review", analysis)
    payload = {
        "workflow": "cooperation_review",
        "standard_contract": standards[0].name,
        "supplier_agreements": [file.name for file in suppliers],
        "analysis": analysis,
    }
    artifacts = [
        _json_artifact("cooperation_review.json", payload),
        WorkflowArtifact(
            _csv_bytes(_cooperation_export_rows(analysis)),
            "cooperation_review.csv",
            "text/csv; charset=utf-8",
        ),
        WorkflowArtifact(
            _analysis_markdown("Cooperation Review", analysis).encode("utf-8"),
            "cooperation_review.md",
            "text/markdown; charset=utf-8",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_factory_certificate(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    del options
    texts: dict[str, str] = {}
    extraction_errors: list[dict[str, str]] = []
    total = len(files)
    for index, file in enumerate(files, 1):
        progress(_file_progress(index - 1, total, 5, 60), f"Extracting {file.name}")
        try:
            text = extract_text_from_pdf(file.data)
            if not text.strip():
                raise ValueError("No readable text found in PDF")
            texts[file.name] = text
        except Exception as exc:  # noqa: BLE001
            extraction_errors.append({"file_name": file.name, "error": _safe_document_error(exc)})
        progress(_file_progress(index, total, 5, 60), f"Processed {file.name}")

    if len(texts) < 2:
        raise ValueError("Factory certificate comparison needs at least two readable documents.")
    progress(65, "Comparing specifications and certificates")
    analysis = _json_mapping(compare_factory_documents(texts))
    if analysis.get("error"):
        raise RuntimeError("The analysis service could not compare the supplied quality documents.")
    analysis = _curated_analysis("factory_certificate", analysis)
    rows = _dict_list(analysis.get("comparisons"))
    payload = {
        "workflow": "factory_certificate",
        "source_files": list(texts),
        "extraction_errors": extraction_errors,
        "analysis": analysis,
    }
    artifacts = [
        _json_artifact("factory_certificate_results.json", payload),
        WorkflowArtifact(_csv_bytes(rows), "factory_certificate_results.csv", "text/csv; charset=utf-8"),
        WorkflowArtifact(
            _xlsx_bytes(rows, "Comparison"),
            "factory_certificate_results.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _run_large_scanner(
    files: list[WorkflowInput], options: Mapping[str, Any], progress: ProgressCallback
) -> WorkflowResult:
    contracts = _files_for_role(files, "contract")
    if len(contracts) != 1:
        raise ValueError("Large scanner requires exactly one contract PDF.")
    contract = contracts[0]
    progress(5, f"Extracting pages from {contract.name}")
    pages = extract_pdf_pages(contract.data)
    readable_pages = [
        page
        for page in pages
        if page.get("extraction_method") != "error" and str(page.get("text", "")).strip()
    ]
    if not readable_pages:
        raise ValueError("No readable text found in the contract PDF.")

    max_chars = _bounded_int(options.get("maxChars", options.get("max_chars")), 18000, 8000, 30000)
    overlap_pages = _bounded_int(options.get("overlapPages", options.get("overlap_pages")), 1, 0, 3)
    complete = _option_bool(options, "isCompleteDocument", "is_complete_document", default=True)
    progress(15, "Building contract chunks")
    chunks = build_contract_chunks(
        readable_pages,
        source_file=contract.name,
        max_chars=max_chars,
        overlap_pages=overlap_pages,
    )
    if not chunks:
        raise ValueError("No analyzable chunks were created from the contract.")

    def on_chunk(done: int, total: int, chunk_result: dict[str, Any]) -> None:
        message = f"Analyzed chunk {done}/{total}"
        if chunk_result.get("error"):
            message = f"Chunk {done}/{total} finished with an error"
        progress(_file_progress(done, total, 20, 75), message)

    chunk_results = analyze_contract_chunks(chunks, progress_callback=on_chunk)
    if chunk_results and all(result.get("error") for result in chunk_results):
        raise RuntimeError("The analysis service could not process the contract sections.")
    progress(80, "Merging contract evidence")
    merged = merge_chunk_findings(chunk_results, chunks)
    report = synthesize_contract_analysis(merged, is_complete_document=complete)
    qa_context = {"version": 1, "chunks": chunks, "merged": merged}
    payload = {
        "workflow": "large_scanner",
        "file_name": contract.name,
        "page_count": len(readable_pages),
        "chunk_results": chunk_results,
        "merged": merged,
        "report": report,
        "settings": {"max_chars": max_chars, "overlap_pages": overlap_pages, "is_complete_document": complete},
        # Later jobs use this durable context with answer_contract_question.
        "qa_context": qa_context,
    }
    artifacts = [
        WorkflowArtifact(
            str(report).encode("utf-8"),
            "large_contract_scanner_report.md",
            "text/markdown; charset=utf-8",
        ),
        _json_artifact("large_contract_scanner_evidence.json", payload),
    ]
    return WorkflowResult(_normalize_json(payload), artifacts)


def _coerce_input_file(value: WorkflowInput | Mapping[str, Any] | Any) -> WorkflowInput:
    if isinstance(value, WorkflowInput):
        return value
    if isinstance(value, Mapping):
        name, role, data, content_type, supplier_hint = (
            value.get("name"),
            value.get("role"),
            value.get("data"),
            value.get("content_type", value.get("contentType")),
            value.get("supplier_hint", value.get("supplierHint", "")),
        )
    else:
        name = getattr(value, "name", None)
        role = getattr(value, "role", None)
        data = getattr(value, "data", None)
        content_type = getattr(value, "content_type", None)
        supplier_hint = getattr(value, "supplier_hint", "")
    if not isinstance(name, str) or not name:
        raise ValueError("Workflow input is missing a filename.")
    if not isinstance(role, str) or not role:
        raise ValueError(f"Workflow input {name} is missing a role.")
    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise TypeError(f"Workflow input {name} must contain bytes.")
    return WorkflowInput(
        PurePosixPath(name.replace("\\", "/")).name,
        role,
        bytes(data),
        content_type,
        str(supplier_hint or ""),
    )


def _normalstunden_pdf_inputs(files: list[WorkflowInput], options: Mapping[str, Any]) -> list[WorkflowInput]:
    pdf_files: list[WorkflowInput] = []
    include_subfolders = _option_bool(options, "includeSubfolders", "include_subfolders", default=True)
    for file in files:
        if file.role == "normalstundenArchive":
            expanded = expand_normalstunden_archive(file.data, get_settings(), include_subfolders)
            pdf_files.extend(
                WorkflowInput(item.name, item.role, item.data, item.content_type, item.supplier_hint) for item in expanded
            )
        elif file.role == "normalstundenPdfs":
            pdf_files.append(file)
    if not pdf_files:
        raise ValueError("Normalstunden workflow requires PDF files or one ZIP archive.")
    return pdf_files


def _files_for_role(files: Iterable[WorkflowInput], role: str) -> list[WorkflowInput]:
    return [file for file in files if file.role == role]


def _load_tender_template(data: bytes) -> tuple[Any, Any, dict[str, tuple[int, int]], tuple[int, int] | None]:
    try:
        workbook = load_workbook(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Could not read tender XLSX template: {exc}") from exc
    worksheet = workbook.worksheets[0]
    field_cells: dict[str, tuple[int, int]] = {}
    source_cell: tuple[int, int] | None = None
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if source_cell is None and isinstance(value, str) and value.strip().lower() == "quelle tenderunterlagen":
                source_cell = (cell.row, cell.column)
        if len(row) >= 5:
            label = row[4].value
            if isinstance(label, str) and label.strip().endswith(":"):
                field_cells[label.strip().rstrip(":").strip()] = (row[4].row, 6)
    return workbook, worksheet, field_cells, source_cell


def _add_market_situation(analysis: dict[str, Any]) -> None:
    extracted = _json_mapping(analysis.get("extracted"))
    if not extracted:
        return
    customer = str(extracted.get("Kundenname", extracted.get("Auftraggeber", ""))).strip()
    project = str(extracted.get("Projekttitel", extracted.get("Leistungsbeschreibung", ""))).strip()
    country = str(extracted.get("Land", "Deutschland")).strip()
    if not customer and not project:
        return
    market = _json_mapping(analyze_market_situation(customer, project, country))
    market.pop("Chancen in %", None)
    extracted.pop("Chancen in %", None)
    for key in ("Vermutliche Wettbewerber", "Letzter Tender", "Split möglich"):
        if key in extracted:
            extracted[key] = market.get(key, "Nicht angegeben")
    analysis["extracted"] = extracted
    analysis["market_situation"] = market


def _tender_display_fields(desired_fields: list[str]) -> list[str]:
    blocked = {"chancen in %", "chance in %", "win probability"}
    if desired_fields:
        return [field for field in desired_fields if field.strip().lower() not in blocked]
    return [
        "customer",
        "project_title",
        "reference_number",
        "procedure",
        "submission_deadline",
        "questions_deadline",
        "contract_start",
        "contract_end",
        "estimated_value",
        "country",
        "language",
        "cpv_codes",
        "notes",
    ]


def _limit_tender_fields(
    analysis: dict[str, Any], display_fields: list[str], *, template_fields: bool
) -> None:
    if template_fields:
        extracted = _json_mapping(analysis.get("extracted"))
        analysis["extracted"] = {
            field: extracted[field] for field in display_fields if field in extracted
        }
        return

    tender_fields = _json_mapping(analysis.get("tender_fields"))
    analysis["tender_fields"] = {
        field: tender_fields[field] for field in display_fields if field in tender_fields
    }


def _merge_tender_results(results: list[dict[str, Any]]) -> dict[str, Any]:
    merged: dict[str, Any] = {"extracted": {}, "field_sources": {}, "german_summary": "", "notes": ""}
    summaries: list[str] = []
    notes: list[str] = []
    for result in results:
        if result.get("status") != "success":
            continue
        analysis = _json_mapping(result.get("analysis"))
        extracted = _json_mapping(analysis.get("extracted", analysis.get("tender_fields", {})))
        for field, value in extracted.items():
            if field.strip().lower() in {"chancen in %", "chance in %", "win probability"}:
                continue
            existing = merged["extracted"].get(field)
            if not existing or existing == "Nicht angegeben":
                merged["extracted"][field] = value
                merged["field_sources"][field] = result["file_name"]
        if analysis.get("german_summary"):
            summaries.append(f"[{result['file_name']}] {analysis['german_summary']}")
        if analysis.get("notes"):
            notes.append(f"[{result['file_name']}] {analysis['notes']}")
    merged["german_summary"] = "\n\n".join(summaries)
    merged["notes"] = "\n\n".join(notes)
    return merged


def _fill_tender_template(
    worksheet: Any,
    field_cells: Mapping[str, tuple[int, int]],
    source_cell: tuple[int, int] | None,
    merged: Mapping[str, Any],
    results: list[dict[str, Any]],
) -> None:
    extracted = _json_mapping(merged.get("extracted"))
    sources = _json_mapping(merged.get("field_sources"))
    for field, (row, column) in field_cells.items():
        cell = worksheet.cell(row=row, column=column)
        if not isinstance(cell, MergedCell):
            cell.value = _excel_value(extracted.get(field, "Nicht angegeben"))
        source = str(sources.get(field, ""))
        source_cell_value = worksheet.cell(row=row, column=column + 1)
        if source and not isinstance(source_cell_value, MergedCell):
            source_cell_value.value = _excel_value(source)
    if merged.get("german_summary"):
        summary_cell = worksheet["C4"]
        if not isinstance(summary_cell, MergedCell):
            summary_cell.value = _excel_value(merged["german_summary"])
    if source_cell:
        cell = worksheet.cell(row=source_cell[0], column=source_cell[1])
        if not isinstance(cell, MergedCell):
            cell.value = _excel_value(
                ", ".join(result["file_name"] for result in results if result.get("status") == "success")
            )


def _add_tender_results_sheet(workbook: Any, results: list[dict[str, Any]]) -> None:
    """Keep arbitrary templates intact while still exporting unmapped tender fields."""
    rows = _tender_export_rows(results)
    worksheet = workbook.create_sheet("Tender Results")
    fields = ["source_file"]
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    worksheet.append(fields)
    for row in rows:
        worksheet.append([_excel_value(row.get(field, "")) for field in fields])


def _tender_export_rows(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        analysis = _json_mapping(result.get("analysis"))
        fields = _json_mapping(analysis.get("extracted", analysis.get("tender_fields", {})))
        rows.append(
            {
                "source_file": result.get("file_name", ""),
                "status": result.get("status", ""),
                "error": result.get("error", ""),
                **fields,
            }
        )
    return rows


def _consolidate_products(products: list[dict[str, Any]], groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    consolidated: list[dict[str, Any]] = []
    for group in groups:
        product_ids = group.get("product_ids", [])
        unique_clients: dict[str, dict[str, Any]] = {}
        for product_id in product_ids if isinstance(product_ids, list) else []:
            if isinstance(product_id, int) and 0 <= product_id < len(products):
                product = products[product_id]
                unique_clients.setdefault(str(product["client_name"]), product)
        if len(unique_clients) < 2:
            continue
        row: dict[str, Any] = {
            "product": group.get("canonical_name", "Unknown Product"),
            "requests": [
                {
                    "client": client,
                    "original_name": product.get("product_name", ""),
                    "quantity": f"{product.get('quantity', '')} {product.get('unit', '')}".strip(),
                    "contract_type": product.get("contract_type", ""),
                }
                for client, product in unique_clients.items()
            ],
        }
        consolidated.append(row)
    return consolidated


def _product_export_rows(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        products = _dict_list(result.get("products"))
        if not products:
            rows.append(
                {
                    "file_name": result.get("file_name"),
                    "status": result.get("status"),
                    "client_name": result.get("client_name", ""),
                    "contract_type": result.get("contract_type", ""),
                    "error": result.get("error", ""),
                }
            )
        for product in products:
            rows.append({**result, **product, "products": None})
    return rows


def _invoice_export_rows(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        products = _dict_list(result.get("products"))
        base = {key: value for key, value in result.items() if key != "products"}
        if products:
            rows.extend({**base, **product} for product in products)
        else:
            rows.append(base)
    return rows


def _normalstunden_export_rows(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        entries = _dict_list(result.get("entries"))
        if entries:
            for entry in entries:
                rows.append(
                    {
                        "file_name": result.get("file_name", ""),
                        "supplier": result.get("supplier", ""),
                        "hours": entry.get("hours_display", entry.get("hours", "")),
                        "hourly_rate": entry.get("hourly_rate_display", entry.get("hourly_rate", "")),
                        "status": result.get("status", ""),
                    }
                )
        else:
            rows.append(
                {
                    "file_name": result.get("file_name", ""),
                    "supplier": result.get("supplier", ""),
                    "hours": result.get("hours_total", ""),
                    "hourly_rate": result.get("hourly_rate_display", ""),
                    "status": result.get("status", ""),
                    "error": result.get("error", ""),
                }
            )
    return rows


def _detailed_export_rows(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        analysis = _json_mapping(result.get("analysis"))
        base = {
            "file_name": result.get("file_name", ""),
            "status": result.get("status", ""),
            "client_name": analysis.get("client_name", ""),
            "contract_type": analysis.get("contract_type", ""),
            "start_date": analysis.get("start_date", ""),
            "end_date": analysis.get("end_date", ""),
            "summary": analysis.get("summary", ""),
            "error": result.get("error", analysis.get("error", "")),
        }
        products = _dict_list(analysis.get("products_services"))
        rows.extend({**base, **product} for product in products) if products else rows.append(base)
    return rows


def _cooperation_export_rows(analysis: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for key, label, title_key in (
        ("deviations", "Deviation", "title"),
        ("risks", "Risk", "title"),
        ("key_clauses", "Clause", "type"),
        ("recommendations", "Recommendation", "action"),
    ):
        for item in _dict_list(analysis.get(key)):
            rows.append({"type": label, "title": item.get(title_key, ""), **item})
    return rows


def _detailed_markdown(results: list[dict[str, Any]]) -> str:
    sections = ["# Detailed Contract Analysis"]
    for result in results:
        sections.append(f"\n## {result.get('file_name', 'Document')}")
        if result.get("status") != "success":
            sections.append(f"\nError: {result.get('error', 'Analysis failed')}")
            continue
        analysis = _json_mapping(result.get("analysis"))
        sections.append(str(analysis.get("summary", "No summary available")))
        if analysis.get("risk_areas"):
            sections.append("\n### Risk Areas\n" + _json_text(analysis["risk_areas"]))
    return "\n".join(sections).strip() + "\n"


def _analysis_markdown(title: str, analysis: Mapping[str, Any]) -> str:
    summary = _json_mapping(analysis.get("summary"))
    summary_text = summary.get("description", analysis.get("summary", "")) if summary else analysis.get("summary", "")
    return f"# {title}\n\n{summary_text}\n\n## Analysis Data\n\n```json\n{_json_text(analysis)}\n```\n"


def _tender_markdown(merged: Mapping[str, Any]) -> str:
    values = _json_mapping(merged.get("extracted"))
    lines = ["# Tender Summary", ""]
    if merged.get("german_summary"):
        lines.extend([str(merged["german_summary"]), ""])
    lines.append("## Extracted Fields")
    lines.extend(f"- **{field}:** {_cell_text(value)}" for field, value in values.items())
    return "\n".join(lines) + "\n"


def _select_fields(value: Mapping[str, Any], *fields: str) -> dict[str, Any]:
    return {field: value[field] for field in fields if field in value}


def _curated_analysis(
    workflow: str,
    analysis: Mapping[str, Any],
    *,
    display_fields: list[str] | None = None,
) -> dict[str, Any]:
    value: dict[str, Any] = {"workflow": workflow}
    if display_fields:
        value["_display_fields"] = display_fields
    if workflow in {"detailed_contract", "tender"}:
        value["results"] = [{"status": "success", "analysis": analysis}]
        result = curate_public_result(workflow, value)
        rows = result.get("results")
        if isinstance(rows, list) and rows and isinstance(rows[0], Mapping):
            return _json_mapping(rows[0].get("analysis"))
        return {}
    value["analysis"] = analysis
    result = curate_public_result(workflow, value)
    return _json_mapping(result.get("analysis"))


def _invoice_fields(extracted: Mapping[str, Any]) -> dict[str, Any]:
    fields = _select_fields(
        extracted,
        "invoice_number",
        "invoice_date",
        "due_date",
        "currency",
        "total_amount",
        "subtotal",
        "tax_amount",
        "tax_rate_percent",
        "payment_terms",
        "po_number",
        "supplier_name",
        "supplier_address",
        "customer_name",
        "customer_address",
        "ship_to",
        "tax_id",
        "contract_type",
        "notes",
    )
    fields["products"] = [
        _select_fields(
            product,
            "product_name",
            "description",
            "quantity",
            "unit",
            "unit_price",
            "line_total",
            "currency",
            "tax_rate_percent",
            "sku_or_part_number",
        )
        for product in _dict_list(extracted.get("products"))
    ]
    return fields


def _summary(results: Iterable[Mapping[str, Any]]) -> dict[str, int]:
    rows = list(results)
    successful = sum(row.get("status") == "success" for row in rows)
    no_match = sum(row.get("status") == "no_match" for row in rows)
    summary = {
        "total": len(rows),
        "successful": successful,
        "failed": sum(row.get("status") == "failed" for row in rows),
    }
    if no_match:
        summary["no_match"] = no_match
    return summary


def _safe_document_error(error: Exception) -> str:
    """Keep persisted job results useful without retaining provider error payloads."""

    message = " ".join(str(error).split())
    if message in {
        "No readable text found in PDF",
        "No readable text found in the contract PDF.",
        "The PDF could not be opened.",
    }:
        return message
    return "Unable to extract or analyze this document."


def _analysis_error_message(analysis: Mapping[str, Any]) -> str:
    """Expose only stable remediation text, never provider error payloads."""
    return ANALYSIS_ERROR_MESSAGES.get(
        str(analysis.get("error", "")),
        "The analysis service could not process this document. Please try again.",
    )


def _json_artifact(name: str, value: Any) -> WorkflowArtifact:
    workflow = str(value.get("workflow", "")) if isinstance(value, Mapping) else ""
    public_value = curate_public_result(workflow, value) if workflow else value
    return WorkflowArtifact(_json_bytes(public_value), name, "application/json")


def _json_bytes(value: Any) -> bytes:
    return _json_text(value).encode("utf-8")


def _json_text(value: Any) -> str:
    return json.dumps(_normalize_json(value), ensure_ascii=False, indent=2, allow_nan=False)


def _csv_bytes(rows: Iterable[Mapping[str, Any]]) -> bytes:
    normalized_rows = [_json_mapping(row) for row in rows]
    fields: list[str] = []
    for row in normalized_rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=fields or ["result"], delimiter=";", extrasaction="ignore")
    writer.writeheader()
    for row in normalized_rows:
        writer.writerow({key: _excel_safe_text(_cell_text(value)) for key, value in row.items()})
    return "\ufeff".encode("utf-8") + output.getvalue().encode("utf-8")


def _xlsx_bytes(rows: Iterable[Mapping[str, Any]], sheet_name: str) -> bytes:
    normalized_rows = [_json_mapping(row) for row in rows]
    fields: list[str] = []
    for row in normalized_rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.append(fields or ["result"])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in normalized_rows:
        worksheet.append([_excel_value(row.get(field, "")) for field in fields] if fields else [""])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in range(1, worksheet.max_column + 1):
        values = (worksheet.cell(row=row, column=column).value for row in range(1, worksheet.max_row + 1))
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            max((len(str(value or "")) for value in values), default=10) + 2,
            48,
        )
    output = BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue()


def _excel_value(value: Any) -> Any:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value if not isinstance(value, float) or math.isfinite(value) else ""
    return _excel_safe_text(_cell_text(value))


def _cell_text(value: Any) -> str:
    value = _normalize_json(value)
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _excel_safe_text(value: str) -> str:
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _normalize_json(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Enum):
        return _normalize_json(value.value)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return {"byte_length": len(value)}
    if hasattr(value, "model_dump") and callable(value.model_dump):
        return _normalize_json(value.model_dump(mode="json"))
    if is_dataclass(value) and not isinstance(value, type):
        return _normalize_json(asdict(value))
    if isinstance(value, Mapping):
        return {str(key): _normalize_json(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_normalize_json(item) for item in value]
    return str(value)


def _json_mapping(value: Any) -> dict[str, Any]:
    normalized = _normalize_json(value)
    if isinstance(normalized, dict):
        return normalized
    if isinstance(normalized, str):
        text = normalized.strip()
        if text.startswith("```"):
            text = text.split("\n", 1)[1] if "\n" in text else ""
            text = text.rsplit("```", 1)[0].strip()
        try:
            parsed = json.loads(text)
            if isinstance(parsed, dict):
                return _normalize_json(parsed)
        except json.JSONDecodeError:
            pass
        return {"raw": normalized}
    return {}


def _as_mapping(value: Any) -> Mapping[str, Any]:
    return _json_mapping(value) if value is not None else {}


def _json_list(value: Any) -> list[Any]:
    normalized = _normalize_json(value)
    return normalized if isinstance(normalized, list) else []


def _dict_list(value: Any) -> list[dict[str, Any]]:
    return [item for item in _json_list(value) if isinstance(item, dict)]


def _job_value(job: Any, name: str, default: Any) -> Any:
    if isinstance(job, Mapping):
        return job.get(name, default)
    return getattr(job, name, default)


def _progress_reporter(callback: ProgressCallback | None) -> ProgressCallback:
    if callback is None:
        return lambda _progress, _message: None

    def report(value: int, message: str) -> None:
        callback(max(0, min(100, int(value))), str(message))

    return report


def _file_progress(index: int, total: int, start: int, end: int) -> int:
    if total <= 0:
        return end
    return start + round((end - start) * index / total)


def _optional_positive_int(value: Any) -> int | None:
    try:
        return int(value) if value is not None and int(value) > 0 else None
    except (TypeError, ValueError):
        return None


def _bounded_int(value: Any, default: int, minimum: int, maximum: int) -> int:
    parsed = _optional_positive_int(value)
    return max(minimum, min(maximum, parsed if parsed is not None else default))


def _option_bool(options: Mapping[str, Any], camel_name: str, snake_name: str, default: bool) -> bool:
    value = options.get(camel_name, options.get(snake_name, default))
    if isinstance(value, str):
        return value.strip().lower() not in {"0", "false", "no", "off"}
    return bool(value)
